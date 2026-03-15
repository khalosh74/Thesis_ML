from __future__ import annotations

from openpyxl.worksheet.worksheet import Worksheet

from Thesis_ML.workbook.structured_sheet_core import fill_simple_structured_sheet
from Thesis_ML.workbook.template_primitives import (
    add_dynamic_named_list,
    add_list_validation,
    add_table,
    col_idx,
    set_widths,
    style_body,
    style_header,
)


def fill_artifact_registry(ws: Worksheet, *, artifact_registry_columns: list[str]) -> int:
    starter = [
        {
            "artifact_id": "",
            "artifact_type": "",
            "run_id": "",
            "status": "",
            "created_at": "",
            "path": "",
            "upstream_artifact_ids": "",
            "config_hash": "",
            "code_ref": "",
            "notes": "Optional workbook mirror of machine registry for audit snapshots.",
        }
    ]
    return fill_simple_structured_sheet(
        ws=ws,
        columns=artifact_registry_columns,
        table_name="WorkbookArtifactRegistryTable",
        title="Workbook Artifact Registry Snapshot",
        width_map={
            "A": 32,
            "B": 20,
            "C": 24,
            "D": 14,
            "E": 24,
            "F": 40,
            "G": 32,
            "H": 24,
            "I": 20,
            "J": 34,
        },
        starter_rows=starter,
    )


def fill_fixed_configs(ws: Worksheet, *, fixed_configs_columns: list[str]) -> int:
    starter = [
        {
            "config_key": "default_target",
            "config_value": "coarse_affect",
            "scope": "global",
            "locked": "No",
            "owner": "Khaled",
            "last_updated": "",
            "notes": "Machine-readable defaults for execution templates.",
        },
        {
            "config_key": "default_model",
            "config_value": "ridge",
            "scope": "global",
            "locked": "No",
            "owner": "Khaled",
            "last_updated": "",
            "notes": "",
        },
    ]
    last = fill_simple_structured_sheet(
        ws=ws,
        columns=fixed_configs_columns,
        table_name="FixedConfigsTable",
        title="Fixed Configurations and Locks",
        width_map={
            "A": 28,
            "B": 28,
            "C": 14,
            "D": 10,
            "E": 16,
            "F": 16,
            "G": 36,
        },
        starter_rows=starter,
    )
    add_list_validation(ws, "=List_YesNo", col_idx(fixed_configs_columns, "locked"), 3, 1000)
    return last


def fill_objectives(ws: Worksheet, *, objectives_columns: list[str]) -> int:
    starter = [
        {
            "objective_id": "OBJ01",
            "objective_text": "Lock target definition under leakage-aware evaluation.",
            "stage": "Stage 1 - Target lock",
            "linked_experiment_id": "E01",
            "primary_metric": "balanced_accuracy",
            "success_criterion": "Decision D01 locked with pre-registered rationale.",
            "status": "Planned",
            "notes": "",
        }
    ]
    last = fill_simple_structured_sheet(
        ws=ws,
        columns=objectives_columns,
        table_name="ObjectivesTable",
        title="Program Objectives",
        width_map={
            "A": 12,
            "B": 44,
            "C": 28,
            "D": 18,
            "E": 20,
            "F": 34,
            "G": 14,
            "H": 30,
        },
        starter_rows=starter,
    )
    add_list_validation(ws, "=List_Stage", col_idx(objectives_columns, "stage"), 3, 1000)
    add_list_validation(
        ws,
        "=List_Experiment_ID",
        col_idx(objectives_columns, "linked_experiment_id"),
        3,
        1000,
    )
    add_list_validation(ws, "=List_Status", col_idx(objectives_columns, "status"), 3, 1000)
    return last


def fill_machine_status(ws: Worksheet, *, machine_status_columns: list[str]) -> int:
    starter = [
        {
            "machine_id": "M01",
            "hostname": "",
            "environment_name": "thesis-dev",
            "python_version": "",
            "gpu": "",
            "status": "Monitoring",
            "last_checked": "",
            "notes": "Track execution environment snapshots used for thesis runs.",
        }
    ]
    last = fill_simple_structured_sheet(
        ws=ws,
        columns=machine_status_columns,
        table_name="MachineStatusTable",
        title="Machine Status and Environment",
        width_map={
            "A": 12,
            "B": 20,
            "C": 22,
            "D": 18,
            "E": 20,
            "F": 14,
            "G": 16,
            "H": 34,
        },
        starter_rows=starter,
    )
    add_list_validation(
        ws,
        "=List_Ethics_Status",
        col_idx(machine_status_columns, "status"),
        3,
        1000,
    )
    return last

def fill_trial_results(ws: Worksheet, *, trial_results_columns: list[str]) -> int:
    last = fill_simple_structured_sheet(
        ws=ws,
        columns=trial_results_columns,
        table_name="TrialResultsTable",
        title="Trial Results (Machine-managed)",
        width_map={
            "A": 22,
            "B": 14,
            "C": 24,
            "D": 14,
            "E": 22,
            "F": 18,
            "G": 34,
            "H": 34,
            "I": 28,
            "J": 30,
            "K": 12,
            "L": 22,
            "M": 34,
            "N": 34,
        },
        starter_rows=None,
    )
    add_list_validation(
        ws, "=List_Experiment_ID", col_idx(trial_results_columns, "experiment_id"), 3, 1000
    )
    add_list_validation(ws, "=List_Status", col_idx(trial_results_columns, "status"), 3, 1000)
    return last


def fill_summary_outputs(ws: Worksheet, *, summary_outputs_columns: list[str]) -> int:
    return fill_simple_structured_sheet(
        ws=ws,
        columns=summary_outputs_columns,
        table_name="SummaryOutputsTable",
        title="Machine Summary Outputs (Best Runs and Patterns)",
        width_map={
            "A": 18,
            "B": 22,
            "C": 20,
            "D": 18,
            "E": 26,
            "F": 14,
            "G": 18,
            "H": 18,
            "I": 14,
            "J": 24,
            "K": 18,
            "L": 20,
            "M": 34,
            "N": 36,
        },
        starter_rows=None,
    )


def fill_experiment_definitions(ws: Worksheet, *, experiment_definitions_columns: list[str]) -> int:
    for idx, header in enumerate(experiment_definitions_columns, start=1):
        ws.cell(1, idx, header)
    style_header(ws, 1, len(experiment_definitions_columns))

    seed_rows = [
        {
            "experiment_id": "E16",
            "enabled": "No",
            "start_section": "dataset_selection",
            "end_section": "evaluation",
            "base_artifact_id": "",
            "target": "coarse_affect",
            "cv": "within_subject_loso_session",
            "model": "ridge",
            "subject": "sub-001",
            "train_subject": "",
            "test_subject": "",
            "filter_task": "",
            "filter_modality": "",
            "reuse_policy": "auto",
            "search_space_id": "",
        },
        {
            "experiment_id": "E17",
            "enabled": "No",
            "start_section": "dataset_selection",
            "end_section": "evaluation",
            "base_artifact_id": "",
            "target": "coarse_affect",
            "cv": "frozen_cross_person_transfer",
            "model": "ridge",
            "subject": "",
            "train_subject": "sub-001",
            "test_subject": "sub-002",
            "filter_task": "",
            "filter_modality": "",
            "reuse_policy": "auto",
            "search_space_id": "",
        },
    ]
    for row_idx, row in enumerate(seed_rows, start=2):
        for col_idx_value, name in enumerate(experiment_definitions_columns, start=1):
            ws.cell(row_idx, col_idx_value, row.get(name, ""))

    last = 81
    style_body(ws, 2, last, 1, len(experiment_definitions_columns))
    add_table(ws, "ExperimentDefinitionsTable", f"A1:O{last}", style="TableStyleMedium2")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:O{last}"

    add_list_validation(
        ws,
        "=List_Experiment_ID",
        col_idx(experiment_definitions_columns, "experiment_id"),
        2,
        1000,
    )
    add_list_validation(
        ws,
        "=List_YesNo",
        col_idx(experiment_definitions_columns, "enabled"),
        2,
        1000,
    )
    add_list_validation(
        ws,
        "=List_Execution_Section",
        col_idx(experiment_definitions_columns, "start_section"),
        2,
        1000,
    )
    add_list_validation(
        ws,
        "=List_Execution_Section",
        col_idx(experiment_definitions_columns, "end_section"),
        2,
        1000,
    )
    add_list_validation(
        ws,
        "=List_Reuse_Policy",
        col_idx(experiment_definitions_columns, "reuse_policy"),
        2,
        1000,
    )
    add_list_validation(
        ws,
        "=List_Search_Space_ID",
        col_idx(experiment_definitions_columns, "search_space_id"),
        2,
        1000,
    )

    set_widths(
        ws,
        {
            "A": 14,
            "B": 10,
            "C": 20,
            "D": 20,
            "E": 28,
            "F": 18,
            "G": 28,
            "H": 14,
            "I": 12,
            "J": 14,
            "K": 14,
            "L": 16,
            "M": 18,
            "N": 18,
            "O": 16,
        },
    )
    return last


def fill_search_spaces(
    ws: Worksheet,
    wb,
    *,
    search_spaces_columns: list[str],
) -> int:
    seed_rows = [
        {
            "search_space_id": "SS01",
            "enabled": "No",
            "optimization_mode": "deterministic_grid",
            "parameter_name": "model",
            "parameter_values": "ridge|logreg|linearsvc",
            "parameter_scope": "parameter",
            "objective_metric": "balanced_accuracy",
            "max_trials": "",
            "notes": "Example model-family search.",
        },
        {
            "search_space_id": "SS01",
            "enabled": "No",
            "optimization_mode": "deterministic_grid",
            "parameter_name": "start_section",
            "parameter_values": "dataset_selection|feature_matrix_load",
            "parameter_scope": "segment",
            "objective_metric": "balanced_accuracy",
            "max_trials": "",
            "notes": "Example section-start search.",
        },
    ]
    last = fill_simple_structured_sheet(
        ws=ws,
        columns=search_spaces_columns,
        table_name="SearchSpacesTable",
        title="Search Space Definitions",
        width_map={
            "A": 16,
            "B": 10,
            "C": 20,
            "D": 24,
            "E": 34,
            "F": 14,
            "G": 20,
            "H": 12,
            "I": 32,
        },
        starter_rows=seed_rows,
    )
    add_list_validation(ws, "=List_YesNo", col_idx(search_spaces_columns, "enabled"), 3, 1000)
    add_list_validation(
        ws,
        "=List_Search_Optimization_Mode",
        col_idx(search_spaces_columns, "optimization_mode"),
        3,
        1000,
    )
    add_list_validation(
        ws,
        "=List_Search_Parameter_Scope",
        col_idx(search_spaces_columns, "parameter_scope"),
        3,
        1000,
    )
    add_dynamic_named_list(
        wb,
        "List_Search_Space_ID",
        ws.title,
        col_idx(search_spaces_columns, "search_space_id"),
        3,
    )
    return last

